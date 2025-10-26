# Run and Plot All Datapoints (Basic + Efficient)
# Authors: Ben Streck (6276247283), Shravya Shashidhar (2763825542), Navya Bhat (5673980946)

import os
import re
import sys
import subprocess
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from openpyxl import Workbook


def sort_key(s):
    """

    :param s:
    :return:
    """
    return [int(text) if text.isdigit() else text.lower() for text in re.split("(\\d+)", s)]


def std_out_break():
    """

    :return:
    """
    print("\n- - - - - - - - - - - - - - - - - - - - - - - - - - - - - - \n")


def extract_file_list(dir_path):
    """

    :param dir_path:
    :return:
    """
    files = [os.path.join(dir_path, f) for f in os.listdir(dir_path) if os.path.isfile(os.path.join(dir_path, f)) and
             ("in" in f)]
    files.sort(key=sort_key)
    return files


def read_outputs(algorithm: str, dir_path):
    """

    :param algorithm:
    :param dir_path:
    :return:
    """
    files = [os.path.join(dir_path, f) for f in os.listdir(dir_path) if os.path.isfile(os.path.join(dir_path, f)) and
             (("out" in f) and (algorithm in f))]
    files.sort(key=sort_key)
    t, m = [], []
    for f in files:
        with open(f, 'r') as file:
            lines = file.readlines()
        t.append(float(lines[3].strip()))
        m.append(float(lines[4].strip()))
    return t, m


def basic_data(dp: list, dir_path):
    """

    :param dp:
    :param dir_path:
    :return:
    """
    for file_in in dp:
        num = re.findall(r"\d+", file_in)[0]
        file_out = dir_path + "out" + num + "_basic.txt"
        subprocess.run([sys.executable, "basic_3.py", file_in, file_out])
        std_out_break()
    time, memory = read_outputs("basic", dir_path)
    return time, memory


def efficient_data(dp: list, dir_path):
    """

    :param dp:
    :param dir_path:
    :return:
    """
    for file_in in dp:
        num = re.findall(r"\d+", file_in)[0]
        file_out = dir_path + "out" + num + "_efficient.txt"
        subprocess.run([sys.executable, "efficient_3.py", file_in, file_out])
        std_out_break()
    time, memory = read_outputs("efficient", dir_path)
    return time, memory


def plot_data(t_b: list, m_b: list, t_e: list, m_e: list):
    """

    :param t_b:
    :param m_b:
    :param t_e:
    :param m_e:
    :return:
    """
    m_plus_n = [16, 64, 128, 256, 384, 512, 768, 1024, 1280, 1536, 2048, 2560, 3072, 3584, 3968]
    with PdfPages('combined_plots.pdf') as pdf:
        # Fig 1 (Memory)
        plt.figure()
        plt.plot(m_plus_n, m_b, marker='o')
        plt.plot(m_plus_n, m_e, marker='o')
        plt.xlabel('M+N')
        plt.ylabel('Memory')
        plt.title('Memory vs Problem Size (M+N)')
        plt.grid(True)
        pdf.savefig()
        plt.close()
        # Fig 2 (Time)
        plt.figure()
        plt.plot(m_plus_n, t_b, marker='o')
        plt.plot(m_plus_n, t_e, marker='o')
        plt.xlabel('M+N')
        plt.ylabel('Time')
        plt.title('Time vs Problem Size (M+N)')
        plt.grid(True)
        pdf.savefig()
        plt.close()

    with PdfPages('basic_only_plots.pdf') as pdf:
        # Fig 1 (Memory)
        plt.figure()
        plt.plot(m_plus_n, m_b, marker='o')
        plt.xlabel('M+N')
        plt.ylabel('Memory')
        plt.title('Memory vs Problem Size (M+N)')
        plt.grid(True)
        pdf.savefig()
        plt.close()
        # Fig 2 (Time)
        plt.figure()
        plt.plot(m_plus_n, t_b, marker='o')
        plt.xlabel('M+N')
        plt.ylabel('Time')
        plt.title('Time vs Problem Size (M+N)')
        plt.grid(True)
        pdf.savefig()
        plt.close()

    with PdfPages('efficient_only_plots.pdf') as pdf:
        # Fig 1 (Memory)
        plt.figure()
        plt.plot(m_plus_n, m_e, marker='o')
        plt.xlabel('M+N')
        plt.ylabel('Memory')
        plt.title('Memory vs Problem Size (M+N)')
        plt.grid(True)
        pdf.savefig()
        plt.close()
        # Fig 2 (Time)
        plt.figure()
        plt.plot(m_plus_n, t_e, marker='o')
        plt.xlabel('M+N')
        plt.ylabel('Time')
        plt.title('Time vs Problem Size (M+N)')
        plt.grid(True)
        pdf.savefig()
        plt.close()
    return


def excel_output(all_data: list):
    """

    :param all_data:
    :return:
    """
    m_plus_n = [16, 64, 128, 256, 384, 512, 768, 1024, 1280, 1536, 2048, 2560, 3072, 3584, 3968]
    sheet_names = ['Basic_Time', 'Efficient_Time', 'Basic_Memory', 'Efficient_Memory']
    wb = Workbook()
    default = wb.active
    wb.remove(default)

    for i in range(len(all_data)):
        ws = wb.create_sheet(title=sheet_names[i])
        ws.cell(row=1, column=1, value="M+N")
        ws.cell(row=1, column=2, value=sheet_names[i])
        for j, (mn, data) in enumerate(zip(m_plus_n, all_data[i]), start=2):
            ws.cell(row=j, column=1, value=mn)
            ws.cell(row=j, column=2, value=data)
    wb.save("Data_Summary.xlsx")
    return


if __name__ == "__main__":
    folder = os.path.join("Datapoints", "")
    datapoints = extract_file_list(folder)
    t1, m1 = basic_data(datapoints, folder)
    t2, m2 = efficient_data(datapoints, folder)
    plot_data(t1, m1, t2, m2)
    excel_output([t1, t2, m1, m2])
